"""Informes de la tienda (balance, Excel, correo) y sugerencias de marketing.

Todo se calcula con SQL y pandas, sin usar el modelo (0 tokens); el modelo solo redacta, de forma opcional, los textos de los avisos sugeridos.
"""
import html
import io
import json
import os
import re
from datetime import date, datetime, timedelta, timezone

import pandas as pd

from contexto_app import _norm, _raiz, _clave
from database import get_connection

CARPETA = os.path.join(os.path.dirname(os.path.abspath(__file__)), "informes")
ESTADOS = ("Pendiente", "Procesado", "Enviado", "Entregado")
VENTA = ("Procesado", "Enviado", "Entregado")
EN_CURSO = ("Procesado", "Enviado")
STOCK_BAJO = 5
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ORDEN_HOJAS = ("Por estado", "Productos vendidos", "Ventas por día", "Por categoría", "Clientes", "Promociones", "Inventario", "Stock por variante", "Órdenes")
# Las fechas de la BD están en UTC; la tienda opera en hora de Colombia (UTC-5, sin horario de verano).
_FECHA = "DATE(o.\"createdAt\" - INTERVAL '5 hours')"

_RE_CORREO = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
_RE_EXCEL = re.compile(r"\b(excel|xlsx|hoja de calculo|planilla)\b")
_RE_INFORME = re.compile(r"\b(informes?|balances?|reportes?)\b")
_RE_ENVIAR = re.compile(r"\b(envia|enviar|envie|enviame|enviamelo|enviale|enviarlo|enviarle|mandar|manda|mandame|mandale|remite|remitir)\b")
_RE_EXCEL_FRASE = re.compile(r"(?:\b(?:en|a|como|con)\s+)?(?:\b(?:un|el)\s+)?(?:\barchivo\s+)?\b(?:excel|xlsx)\b", re.IGNORECASE)

_DIMENSIONES = (
    (r"\b(stock|inventario|existencias?)\b", ("Inventario", "Stock por variante")),
    (r"\b(por\s+(fechas?|dias?)|diari[oa]s?|fechas?)\b", ("Ventas por día",)),
    (r"\bcategori[ao]s?\b", ("Por categoría",)),
    (r"\bclientes?\b", ("Clientes",)),
    (r"\b(productos?|vendid[oa]s?|articulos?)\b", ("Productos vendidos",)),
    (r"\b(promo\w*|codigos?|descuentos?|cupon\w*)\b", ("Promociones",)),
    (r"\b(estados?|pendientes?|entregad[oa]s?|en\s+curso|enviad[oa]s?|procesad[oa]s?)\b", ("Por estado",)),
    (r"\b(ordenes|pedidos?|detalle)\b", ("Órdenes",)),
)
# Palabras que forman parte de pedir un informe; si "excel" va acompañado de otras, es una consulta con filtros.
_VOCAB = {_raiz(p) for p in (
    "excel", "xlsx", "hoja", "calculo", "planilla", "archivo", "informe", "balance", "reporte", "semanal", "semana", "mensual", "mes", "ano",
    "hoy", "ayer", "pasada", "anterior", "ultimos", "ultimo", "esta", "este", "dias", "dia", "desde", "hasta", "fecha", "fechas", "diario",
    "stock", "inventario", "existencias", "categoria", "categorias", "cliente", "clientes", "producto", "productos", "vendido", "vendidos",
    "venta", "ventas", "promocion", "promociones", "codigo", "codigos", "descuento", "descuentos", "estado", "estados", "pendiente",
    "pendientes", "entregado", "entregados", "curso", "enviado", "enviados", "procesado", "procesados", "orden", "ordenes", "pedido",
    "pedidos", "detalle", "envia", "enviar", "enviame", "manda", "mandame", "correo", "email", "mail", "segun", "pida", "le", "lo", "se",
    "cuanto", "cuantos", "cuantas", "total", "totales")}


# ---------------------------------------------------------------- Detección de la petición


def hoy_colombia():
    return (datetime.now(timezone.utc) - timedelta(hours=5)).date()


def _fecha(d, m, a, hoy):
    anio = int(a) if a else hoy.year
    if anio < 100:
        anio += 2000
    return date(anio, int(m), int(d))


def _etiqueta(ini, fin):
    return f"{ini:%d/%m/%Y}" if ini == fin else f"{ini:%d/%m/%Y} al {fin:%d/%m/%Y}"


def parsear_periodo(texto, hoy=None):
    """(inicio, fin, etiqueta) según el texto; sin indicación, los últimos 7 días."""
    hoy = hoy or hoy_colombia()
    t = _norm(texto)
    fechas = []
    for a, m, d in re.findall(r"\b(\d{4})-(\d{2})-(\d{2})\b", t):
        try:
            fechas.append(date(int(a), int(m), int(d)))
        except ValueError:
            raise ValueError(f"La fecha {d}/{m}/{a} no es válida.") from None
    sin_iso = re.sub(r"\b\d{4}-\d{2}-\d{2}\b", " ", t)
    for d, m, a in re.findall(r"\b(\d{1,2})[/-](\d{1,2})(?:[/-](\d{2,4}))?\b", sin_iso):
        try:
            fechas.append(_fecha(d, m, a, hoy))
        except ValueError:
            raise ValueError(f"La fecha {d}/{m}{'/' + a if a else ''} no es válida.") from None
    if len(fechas) >= 2:
        ini, fin = sorted(fechas[:2])
        return ini, fin, _etiqueta(ini, fin)
    if fechas:
        ini = fechas[0]
        fin = hoy if re.search(r"\bdesde\b", t) and ini <= hoy else ini
        return ini, fin, _etiqueta(ini, fin)

    lunes = hoy - timedelta(days=hoy.weekday())
    n = re.search(r"ultim[oa]s?\s+(\d{1,3})\s+dias?", t)
    if n:
        dias = max(1, int(n.group(1)))
        return hoy - timedelta(days=dias - 1), hoy, f"últimos {dias} días"
    if re.search(r"\bayer\b", t):
        ayer = hoy - timedelta(days=1)
        return ayer, ayer, "ayer"
    if re.search(r"\bhoy\b", t):
        return hoy, hoy, "hoy"
    if re.search(r"semana (pasada|anterior)", t):
        return lunes - timedelta(days=7), lunes - timedelta(days=1), "semana pasada"
    if re.search(r"\besta semana\b", t):
        return lunes, hoy, "esta semana"
    if re.search(r"mes (pasado|anterior)", t):
        fin = hoy.replace(day=1) - timedelta(days=1)
        return fin.replace(day=1), fin, "mes pasado"
    if re.search(r"\b(este mes|mensual|del mes)\b", t):
        return hoy.replace(day=1), hoy, "este mes"
    if re.search(r"\bultimo mes\b", t):
        return hoy - timedelta(days=29), hoy, "últimos 30 días"
    if re.search(r"\b(este ano|anual)\b", t):
        return date(hoy.year, 1, 1), hoy, "este año"
    return hoy - timedelta(days=6), hoy, "últimos 7 días"


def seleccion_hojas(texto):
    """Hojas del informe que pide el texto (por stock, fechas, clientes...); sin indicación, todas."""
    t = _RE_CORREO.sub(" ", _norm(texto))
    elegidas = set()
    for patron, hojas in _DIMENSIONES:
        if re.search(patron, t):
            elegidas.update(hojas)
    if not elegidas and re.search(r"\bventas?\b", t):
        elegidas.update(("Productos vendidos", "Ventas por día"))
    return tuple(h for h in ORDEN_HOJAS if h in elegidas) or ORDEN_HOJAS


def detectar_informe(texto):
    """Dict con período, hojas y correo si la petición es un balance/informe o un Excel genérico; None si no."""
    t = _norm(texto)
    sin_correo = _RE_CORREO.sub(" ", t)
    informe = bool(_RE_INFORME.search(sin_correo)) or bool(
        re.search(r"\bresumen\b", sin_correo) and re.search(r"\b(semanal|semana|mensual|mes)\b", sin_correo))
    excel = bool(_RE_EXCEL.search(sin_correo))
    if not informe and not excel:
        return None
    if not informe and {w for w in _clave(sin_correo) if not w.isdigit()} - _VOCAB:
        return None
    correos = _RE_CORREO.findall(texto)
    correo = correos[0] if correos else None
    try:
        periodo = parsear_periodo(_RE_CORREO.sub(" ", texto))
    except ValueError as e:
        return {"error": str(e)}
    return {
        "periodo": periodo,
        "hojas": seleccion_hojas(texto),
        "correo": correo,
        "falta_correo": not correo and bool(_RE_ENVIAR.search(sin_correo)),
    }


def detectar_sugerencia(texto):
    """'avisos', 'codigos' o 'ambos' si piden sugerencias de marketing; None si no."""
    t = _norm(texto)
    if not re.search(r"\b(sugier\w*|sugerenc\w*|recomiend\w*|recomendac\w*|ideas?|propon\w*|propuestas?)\b", t):
        return None
    avisos = re.search(r"\b(avisos?|anuncios?|publicac\w*|ofertas?|banners?|colecc\w*)\b", t)
    codigos = re.search(r"\b(codigos?|cupon\w*|promoc\w*|descuentos?)\b", t)
    if avisos and not codigos:
        return "avisos"
    if codigos and not avisos:
        return "codigos"
    if avisos or codigos or re.search(r"\b(vender|marketing|campan\w*)\b", t):
        return "ambos"
    return None


def quitar_excel(texto):
    """Quita 'en excel', 'a excel'... para que el resto se trate como una consulta normal."""
    texto = re.sub(r"\b(exporta\w*|descarga\w*|genera\w*)\b", " ", texto, flags=re.IGNORECASE)
    texto = _RE_EXCEL_FRASE.sub(" ", texto)
    texto = re.sub(r"\b(dame|quiero|necesito)\s+(de|el|la)\b", r"\1", texto, flags=re.IGNORECASE)
    return re.sub(r"\s{2,}", " ", texto).strip()


# ---------------------------------------------------------------- Datos del balance


def _df(cur, sql, params=()):
    cur.execute(sql, params)
    return pd.DataFrame(cur.fetchall(), columns=[c[0] for c in cur.description])


def _estado_stock(stock):
    return "Sin stock" if stock <= 0 else ("Stock bajo" if stock <= STOCK_BAJO else "Disponible")


def _estado_codigo(r, ahora):
    if not r["activo"]:
        return "Inactivo"
    if not pd.isna(r["vence"]) and r["vence"] < ahora:
        return "Expirado"
    if not pd.isna(r["usos_max"]) and r["usos"] >= r["usos_max"]:
        return "Agotado"
    return "Activo"


def datos_balance(inicio, fin, cur=None, etiqueta=None):
    """Balance del período [inicio, fin]: resumen, hojas (DataFrames) y filas del resumen para Excel/correo."""
    if cur is None:
        conn = get_connection()
        try:
            return datos_balance(inicio, fin, conn.cursor(), etiqueta)
        finally:
            conn.close()
    etiqueta = etiqueta or _etiqueta(inicio, fin)
    p = (inicio, fin)
    ordenes = _df(cur, f"""
        SELECT o.id AS orden, {_FECHA} AS fecha, c.name AS cliente, c.email, c."isVIP" AS vip, o.status AS estado,
               o.total, o.discount AS descuento, COALESCE(o."promoCode", '') AS codigo
        FROM orders o JOIN customers c ON c.id = o."customerId"
        WHERE {_FECHA} BETWEEN %s AND %s ORDER BY o."createdAt" DESC""", p)
    items = _df(cur, f"""
        SELECT o.id AS orden, o.status AS estado, {_FECHA} AS fecha, p.id AS producto_id, p.name AS producto, p.category AS categoria,
               oi.quantity AS unidades, oi.price AS precio, oi.quantity * oi.price AS subtotal
        FROM order_items oi JOIN orders o ON o.id = oi."orderId" JOIN products p ON p.id = oi."productId"
        WHERE {_FECHA} BETWEEN %s AND %s""", p)
    cur.execute("""SELECT COUNT(*) FROM customers WHERE DATE("createdAt" - INTERVAL '5 hours') BETWEEN %s AND %s""", p)
    nuevos = int(cur.fetchone()[0])
    inventario = _df(cur, """
        SELECT p.id, p.sku, p.name AS producto, p.category AS categoria, p.genero, p.size AS talla, p.color, p.price AS precio,
               p."costPrice" AS costo, p.stock,
               (SELECT COALESCE(SUM(v.stock), 0) FROM product_variants v WHERE v."productId" = p.id) AS stock_variantes
        FROM products p ORDER BY p.name""")
    variantes = _df(cur, """
        SELECT p.name AS producto, v.size AS talla, v.color, v.stock, v.sku
        FROM product_variants v JOIN products p ON p.id = v."productId" ORDER BY p.name, v.size""")
    codigos = _df(cur, """
        SELECT code AS codigo, "discountPct" AS descuento_pct, active AS activo, "maxUses" AS usos_max, "usedCount" AS usos, "expiresAt" AS vence
        FROM promo_codes ORDER BY "createdAt" DESC""")

    ventas = ordenes[ordenes["estado"].isin(VENTA)]
    items_v = items[items["estado"].isin(VENTA)]

    por_estado = pd.DataFrame({"estado": list(ESTADOS)})
    por_estado["ordenes"] = [int((ordenes["estado"] == e).sum()) for e in ESTADOS]
    por_estado["unidades"] = [int(items.loc[items["estado"] == e, "unidades"].sum()) for e in ESTADOS]
    por_estado["total"] = [float(ordenes.loc[ordenes["estado"] == e, "total"].sum()) for e in ESTADOS]

    columnas_prod = ["producto_id", "producto", "categoria", "unidades", "ventas_brutas"]
    if len(items_v):
        productos = (items_v.groupby(["producto_id", "producto", "categoria"], as_index=False)
                     .agg(unidades=("unidades", "sum"), ventas_brutas=("subtotal", "sum"))
                     .sort_values(["unidades", "ventas_brutas"], ascending=False))
        por_categoria = (items_v.groupby("categoria", as_index=False)
                         .agg(unidades=("unidades", "sum"), ventas_brutas=("subtotal", "sum"))
                         .sort_values("unidades", ascending=False))
    else:
        productos = pd.DataFrame(columns=columnas_prod)
        por_categoria = pd.DataFrame(columns=["categoria", "unidades", "ventas_brutas"])

    dias = list(pd.date_range(inicio, fin).date)
    por_dia = pd.DataFrame({"fecha": dias})
    por_dia["ordenes"] = [int((ventas["fecha"] == d).sum()) for d in dias]
    por_dia["unidades"] = [int(items_v.loc[items_v["fecha"] == d, "unidades"].sum()) for d in dias]
    por_dia["ingresos"] = [float(ventas.loc[ventas["fecha"] == d, "total"].sum()) for d in dias]

    if len(ordenes):
        g = ordenes.assign(es_venta=ordenes["estado"].isin(VENTA), gasto=ordenes["total"].where(ordenes["estado"].isin(VENTA), 0.0))
        clientes = (g.groupby(["cliente", "email", "vip"], as_index=False)
                    .agg(ordenes=("orden", "count"), compras=("es_venta", "sum"), gastado=("gasto", "sum"))
                    .sort_values("gastado", ascending=False))
        uso = (ordenes[ordenes["codigo"] != ""].groupby("codigo")
               .agg(usos_periodo=("orden", "count"), descuento_periodo=("descuento", "sum")))
    else:
        clientes = pd.DataFrame(columns=["cliente", "email", "vip", "ordenes", "compras", "gastado"])
        uso = pd.DataFrame(columns=["usos_periodo", "descuento_periodo"])
    if len(codigos):
        codigos = codigos.merge(uso, how="left", left_on="codigo", right_index=True)
        codigos[["usos_periodo", "descuento_periodo"]] = codigos[["usos_periodo", "descuento_periodo"]].apply(pd.to_numeric, errors="coerce").fillna(0)
        ahora = datetime.now(timezone.utc).replace(tzinfo=None)
        codigos["estado"] = codigos.apply(lambda r: _estado_codigo(r, ahora), axis=1)

    if len(inventario):
        inventario["estado_stock"] = inventario["stock"].apply(_estado_stock)
        inventario["valor_inventario"] = inventario["precio"] * inventario["stock"]
    else:
        inventario["estado_stock"] = []
        inventario["valor_inventario"] = []

    n = len(ordenes)
    cuenta = {e: int((ordenes["estado"] == e).sum()) for e in ESTADOS}
    en_curso = sum(cuenta[e] for e in EN_CURSO)
    unidades_vendidas = int(items_v["unidades"].sum())
    ingresos = float(ventas["total"].sum())
    resumen = {
        "ordenes": n, "pendientes": cuenta["Pendiente"], "procesadas": cuenta["Procesado"], "enviadas": cuenta["Enviado"],
        "entregadas": cuenta["Entregado"], "en_curso": en_curso, "unidades_vendidas": unidades_vendidas, "ingresos": ingresos,
        "descuentos": float(ventas["descuento"].sum()), "ticket": ingresos / len(ventas) if len(ventas) else 0.0,
        "con_codigo": int((ordenes["codigo"] != "").sum()), "clientes_nuevos": nuevos,
        "sin_stock": int((inventario["stock"] <= 0).sum()) if len(inventario) else 0,
        "stock_bajo": int(((inventario["stock"] > 0) & (inventario["stock"] <= STOCK_BAJO)).sum()) if len(inventario) else 0,
    }
    filas = [
        ("Período", etiqueta),
        ("Órdenes en el período", resumen["ordenes"]),
        ("Pendientes", resumen["pendientes"]),
        ("En curso (procesadas + enviadas)", resumen["en_curso"]),
        ("   · Procesadas", resumen["procesadas"]),
        ("   · Enviadas", resumen["enviadas"]),
        ("Entregadas", resumen["entregadas"]),
        ("Unidades vendidas (órdenes procesadas, enviadas o entregadas)", resumen["unidades_vendidas"]),
        ("Unidades en órdenes pendientes", int(items.loc[items["estado"] == "Pendiente", "unidades"].sum())),
        ("Ingresos (después de descuentos)", resumen["ingresos"]),
        ("Descuentos aplicados", resumen["descuentos"]),
        ("Ticket promedio", resumen["ticket"]),
        ("Órdenes con código promocional", resumen["con_codigo"]),
        ("Clientes nuevos", resumen["clientes_nuevos"]),
        ("Productos sin stock (hoy)", resumen["sin_stock"]),
        (f"Productos con stock bajo (hoy, hasta {STOCK_BAJO})", resumen["stock_bajo"]),
    ]
    hojas = {
        "Por estado": por_estado, "Productos vendidos": productos, "Ventas por día": por_dia, "Por categoría": por_categoria,
        "Clientes": clientes, "Promociones": codigos, "Inventario": inventario, "Stock por variante": variantes,
        "Órdenes": ordenes[["orden", "fecha", "cliente", "estado", "total", "descuento", "codigo"]],
    }
    return {"inicio": inicio, "fin": fin, "etiqueta": etiqueta, "resumen": resumen, "filas": filas, "hojas": hojas}


# ---------------------------------------------------------------- Excel y correo

_MONEDA = ("ingresos", "total", "precio", "costo", "valor", "descuento", "gastado", "ventas_brutas", "subtotal")


def _sanear(df):
    """Evita inyección de fórmulas en Excel y las zonas horarias, que Excel no admite."""
    df = df.copy()
    for c in df.columns:
        if str(df[c].dtype).startswith("datetime64") and getattr(df[c].dt, "tz", None) is not None:
            df[c] = df[c].dt.tz_localize(None)
        elif df[c].dtype == object:
            df[c] = df[c].map(lambda v: "'" + v if isinstance(v, str) and v[:1] in ("=", "+", "-", "@") else v)
    return df


def _escribir_hoja(xw, nombre, df):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    df = _sanear(df)
    df.to_excel(xw, sheet_name=nombre, index=False)
    ws = xw.sheets[nombre]
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="B0001B")
        celda.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    if len(df):
        ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(df.columns, start=1):
        ancho = max([len(str(col))] + [len(str(v)) for v in df[col].head(200)]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(ancho, 60)
        nombre_col = str(col).lower()
        if any(k in nombre_col for k in _MONEDA) and "pct" not in nombre_col:
            for fila in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                fila[0].number_format = '"$"#,##0'
    return ws


def construir_excel(informe, hojas=ORDEN_HOJAS):
    """Bytes de un .xlsx con la hoja Resumen y las hojas pedidas."""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        ws = _escribir_hoja(xw, "Resumen", pd.DataFrame(informe["filas"], columns=["Concepto", "Valor"]))
        ws.column_dimensions["A"].width = 62
        ws.column_dimensions["B"].width = 26
        for fila in ws.iter_rows(min_row=2, max_col=2):
            if any(k in str(fila[0].value) for k in ("Ingresos", "Descuentos aplicados", "Ticket")):
                fila[1].number_format = '"$"#,##0'
        for h in hojas:
            _escribir_hoja(xw, h, informe["hojas"][h])
    return buf.getvalue()


def excel_de_resultado(filas, columnas, hoja="Consulta"):
    """Bytes de un .xlsx con el resultado de cualquier consulta."""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        _escribir_hoja(xw, hoja, pd.DataFrame(list(filas), columns=list(columnas)))
    return buf.getvalue()


def guardar_excel(datos, inicio, fin):
    """Guarda el Excel en la carpeta informes/ y devuelve su ruta."""
    os.makedirs(CARPETA, exist_ok=True)
    ruta = os.path.join(CARPETA, f"balance_{inicio:%Y-%m-%d}_{fin:%Y-%m-%d}_{datetime.now():%H%M%S}.xlsx")
    with open(ruta, "wb") as f:
        f.write(datos)
    return ruta


def _valor_texto(concepto, valor):
    if isinstance(valor, float):
        return f"${valor:,.0f}".replace(",", ".")
    return str(valor)


def cuerpo_correo(informe, hojas):
    """(texto, html) del correo con el resumen del balance; el detalle va en el Excel adjunto."""
    filas = [(c, _valor_texto(c, v)) for c, v in informe["filas"]]
    texto = "\n".join([f"BALANCE CROVN - {informe['etiqueta']}", ""] + [f"{c}: {v}" for c, v in filas] +
                      ["", f"El Excel adjunto incluye las hojas: Resumen, {', '.join(hojas)}."])
    celdas = "".join(
        f"<tr><td style='padding:6px 12px;border-bottom:1px solid #eee'>{html.escape(c)}</td>"
        f"<td style='padding:6px 12px;border-bottom:1px solid #eee;text-align:right'><b>{html.escape(v)}</b></td></tr>"
        for c, v in filas)
    cuerpo = (
        "<div style=\"font-family:-apple-system,Segoe UI,Roboto,sans-serif;max-width:640px;margin:0 auto;color:#1a1a1a\">"
        "<div style='background:#1a1a1a;color:#fff;padding:24px;text-align:center'>"
        "<h1 style='margin:0;font-size:24px;letter-spacing:2px'>CROVN<span style='color:#e63946'>.</span></h1>"
        f"<p style='margin:6px 0 0;color:#aaa;font-size:13px'>Balance · {html.escape(informe['etiqueta'])}</p></div>"
        f"<table style='width:100%;border-collapse:collapse;font-size:14px;margin:16px 0'>{celdas}</table>"
        f"<p style='font-size:12px;color:#888'>El Excel adjunto incluye las hojas: Resumen, {html.escape(', '.join(hojas))}.</p></div>")
    return texto, cuerpo


# ---------------------------------------------------------------- Sugerencias de marketing


def _aviso(titulo, subtitulo, tipo, producto, descuento, motivo):
    return {"clase": "aviso", "titulo": titulo, "subtitulo": subtitulo, "tipo": tipo, "producto_id": int(producto["id"]),
            "producto": producto["name"], "descuento": descuento, "motivo": motivo}


def _codigo(codigo, descuento, usos, vence, motivo):
    return {"clase": "codigo", "codigo": codigo, "descuento": descuento, "usos_max": usos, "vence": vence, "motivo": motivo}


def ideas_marketing(tipo="ambos", cur=None, hoy=None):
    """Avisos y códigos promocionales sugeridos a partir de las ventas, el inventario y los clientes de los últimos 30 días."""
    if cur is None:
        conn = get_connection()
        try:
            return ideas_marketing(tipo, conn.cursor(), hoy)
        finally:
            conn.close()
    hoy = hoy or hoy_colombia()
    desde = hoy - timedelta(days=30)
    ideas = []
    if tipo in ("avisos", "ambos"):
        top = _df(cur, f"""
            SELECT p.id, p.name, SUM(oi.quantity) AS unidades FROM order_items oi
            JOIN orders o ON o.id = oi."orderId" JOIN products p ON p.id = oi."productId"
            WHERE o.status IN %s AND {_FECHA} >= %s GROUP BY p.id, p.name ORDER BY unidades DESC LIMIT 1""", (VENTA, desde))
        quieto = _df(cur, f"""
            SELECT p.id, p.name, p.stock FROM products p
            WHERE p.stock >= 5 AND NOT EXISTS (
                SELECT 1 FROM order_items oi JOIN orders o ON o.id = oi."orderId"
                WHERE oi."productId" = p.id AND o.status IN %s AND {_FECHA} >= %s)
            ORDER BY p.stock DESC LIMIT 1""", (VENTA, desde))
        bajo = _df(cur, "SELECT id, name, stock FROM products WHERE stock BETWEEN 1 AND %s ORDER BY stock LIMIT 1", (STOCK_BAJO,))
        nuevo = _df(cur, """SELECT id, name FROM products ORDER BY "createdAt" DESC LIMIT 1""")
        for r in top.to_dict("records"):
            ideas.append(_aviso(f"Lo más vendido: {r['name']}", f"{int(r['unidades'])} unidades vendidas en 30 días", "noticia", r, 0,
                                "Es tu producto con más ventas de los últimos 30 días."))
        for r in quieto.to_dict("records"):
            d = 20 if r["stock"] >= 15 else 15
            ideas.append(_aviso(f"Oferta en {r['name']}", f"{d}% de descuento por tiempo limitado", "oferta", r, d,
                                f"Tienes {int(r['stock'])} unidades y no se ha vendido ninguna en 30 días."))
        for r in bajo.to_dict("records"):
            ideas.append(_aviso(f"¡Últimas unidades de {r['name']}!", f"Quedan {int(r['stock'])}, no te quedes sin la tuya", "noticia", r, 0,
                                "Stock bajo: genera urgencia de compra."))
        if len(ideas) < 2:
            for r in nuevo.to_dict("records"):
                ideas.append(_aviso(f"Nuevo en CROVN: {r['name']}", "Descúbrelo ya en la tienda", "coleccion", r, 0,
                                    "Es el último producto agregado al catálogo."))
    if tipo in ("codigos", "ambos"):
        cur.execute("""
            SELECT COUNT(*), COUNT(*) FILTER (WHERE "isVIP"),
                   COUNT(*) FILTER (WHERE NOT EXISTS (SELECT 1 FROM orders o WHERE o."customerId" = c.id
                                    AND DATE(o."createdAt" - INTERVAL '5 hours') >= %s))
            FROM customers c""", (desde,))
        total, vip, inactivos = (int(x) for x in cur.fetchone())
        cur.execute("SELECT UPPER(code) FROM promo_codes")
        existentes = {r[0] for r in cur.fetchall()}

        def unico(base):
            codigo = base if base not in existentes else f"{base}{hoy:%d%m}"
            existentes.add(codigo)
            return codigo

        codigos = [_codigo(unico("BIENVENIDO10"), 10, 100, hoy + timedelta(days=30),
                           f"Atrae clientes nuevos (hoy hay {total} registrados).")]
        if inactivos:
            codigos.append(_codigo(unico("VUELVE15"), 15, max(inactivos, 10), hoy + timedelta(days=14),
                                   f"{inactivos} cliente(s) no compran desde hace 30 días o más."))
        if vip:
            codigos.append(_codigo(unico("VIP20"), 20, max(vip * 3, 5), hoy + timedelta(days=30),
                                   f"Premia a tus {vip} cliente(s) VIP."))
        if len(codigos) < 2:
            codigos.append(_codigo(unico("SEMANA10"), 10, 50, hoy + timedelta(days=7), "Impulsa las ventas de la semana."))
        ideas.extend(codigos)
    for n, idea in enumerate(ideas):
        idea["id"] = f"{datetime.now():%H%M%S}{n}"
    return ideas


def mejorar_textos(ideas, llamar):
    """Pide al modelo textos más atractivos para los avisos; ante cualquier fallo conserva los originales."""
    avisos = [x for x in ideas if x["clase"] == "aviso"]
    if not avisos:
        return ideas
    lineas = "\n".join(f"{n}: {x['titulo']} | {x['subtitulo']} | {x['tipo']}" for n, x in enumerate(avisos))
    prompt = ("Eres el redactor de CROVN, tienda de camisetas y ropa urbana en Colombia. Reescribe el título (máx. 6 palabras) y el subtítulo "
              "(máx. 12 palabras) de cada aviso con tono cercano, sin emojis y sin cambiar nombres de productos, cifras ni porcentajes. "
              'Responde solo JSON: {"avisos":[{"n":0,"titulo":"...","subtitulo":"..."}]}\n' + lineas)
    try:
        datos = json.loads(llamar(prompt, json_mode=True, max_tokens=300))
        for item in datos.get("avisos", []):
            aviso = avisos[int(item["n"])]
            titulo, sub = str(item["titulo"]).strip(), str(item["subtitulo"]).strip()
            original = f"{aviso['titulo']} {aviso['subtitulo']}"
            cifras_ok = set(re.findall(r"\d+", original)) <= set(re.findall(r"\d+", f"{titulo} {sub}"))
            if (titulo and sub and len(titulo) <= 70 and len(sub) <= 120 and "\n" not in titulo + sub
                    and aviso["producto"].lower() in f"{titulo} {sub}".lower() and cifras_ok):
                aviso["titulo"], aviso["subtitulo"] = titulo, sub
    except Exception:
        pass
    return ideas


def texto_ideas(ideas):
    """Resumen en Markdown de las sugerencias (queda en el historial del chat)."""
    avisos = [x for x in ideas if x["clase"] == "aviso"]
    codigos = [x for x in ideas if x["clase"] == "codigo"]
    partes = []
    if avisos:
        partes.append("**Avisos sugeridos**")
        for x in avisos:
            dto = f", {x['descuento']}% de descuento" if x["descuento"] else ""
            partes.append(f"- **{x['titulo']}** ({x['tipo']}{dto}) — {x['subtitulo']}. _{x['motivo']}_")
    if codigos:
        partes.append("**Códigos promocionales sugeridos**")
        for x in codigos:
            partes.append(f"- **{x['codigo']}** — {x['descuento']}% hasta el {x['vence']:%d/%m/%Y}, máximo {x['usos_max']} usos. _{x['motivo']}_")
    return "\n\n".join(partes) or "No encontré datos suficientes para sugerir."


def crear_publicacion(idea, cur=None):
    """Crea el aviso sugerido en la tabla publications; (ok, mensaje)."""
    if cur is None:
        conn = get_connection()
        try:
            resultado = crear_publicacion(idea, conn.cursor())
            conn.commit()
            return resultado
        finally:
            conn.close()
    cur.execute("SELECT 1 FROM publications WHERE LOWER(title) = LOWER(%s) LIMIT 1", (idea["titulo"],))
    if cur.fetchone():
        return False, f"Ya existe una publicación con el título «{idea['titulo']}»."
    cur.execute(
        """INSERT INTO publications (title, subtitle, type, active, "productId", "discountPct", "updatedAt")
           VALUES (%s, %s, %s, true, %s, %s, NOW())""",
        (idea["titulo"], idea["subtitulo"], idea["tipo"], idea["producto_id"], idea["descuento"]))
    return True, f"Aviso «{idea['titulo']}» creado y activo."


def crear_codigo(idea, cur=None):
    """Crea el código promocional sugerido en la tabla promo_codes; (ok, mensaje)."""
    if cur is None:
        conn = get_connection()
        try:
            resultado = crear_codigo(idea, conn.cursor())
            conn.commit()
            return resultado
        finally:
            conn.close()
    cur.execute("SELECT 1 FROM promo_codes WHERE UPPER(code) = UPPER(%s) LIMIT 1", (idea["codigo"],))
    if cur.fetchone():
        return False, f"El código {idea['codigo']} ya existe."
    cur.execute(
        """INSERT INTO promo_codes (code, "discountPct", active, "maxUses", "usedCount", "expiresAt", "updatedAt")
           VALUES (%s, %s, true, %s, 0, %s, NOW())""",
        (idea["codigo"], idea["descuento"], idea["usos_max"], idea["vence"]))
    return True, f"Código {idea['codigo']} creado ({idea['descuento']}% hasta el {idea['vence']:%d/%m/%Y})."
